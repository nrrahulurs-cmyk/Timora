import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml

# -------------------------------------------------------------------
# EXACT COLORS FROM Timora_Project_Presentation.pdf
# -------------------------------------------------------------------
COLOR_BG = RGBColor(251, 250, 246)          # Warm Ivory #FBFAF6
COLOR_DARK = RGBColor(32, 39, 34)           # Charcoal Black #202722
COLOR_MUTED = RGBColor(97, 112, 106)        # Sage Slate #61706A
COLOR_RED = RGBColor(183, 68, 45)           # Terracotta Brick Red #B7442D
COLOR_PILL_BG = RGBColor(220, 236, 224)     # Soft Sage Green #DCECE0
COLOR_PILL_TEXT = RGBColor(23, 36, 31)      # Deep Forest Dark #17241F
COLOR_BORDER = RGBColor(221, 228, 220)      # Border Grey #DDE4DC
COLOR_SHADOW = RGBColor(231, 234, 231)      # Offset Shadow Fill #E7EAE7
COLOR_CARD_BG = RGBColor(241, 244, 239)     # Soft Card Fill #F1F4EF
COLOR_BANNER_BG = RGBColor(23, 36, 31)      # Deep Forest Charcoal #17241F
COLOR_WHITE = RGBColor(255, 255, 255)

FONT_FAMILY = "Arial"

def pt_to_in(pt):
    return Inches(pt / 72.0)

def set_slide_transition(slide, transition_type="fade", speed="med", direction="l"):
    """Injects high-fidelity PowerPoint transition tags into OpenXML."""
    for child in list(slide._element):
        if child.tag.endswith("transition"):
            slide._element.remove(child)
            
    if transition_type == "push":
        xml = f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="{speed}"><p:push dir="{direction}"/></p:transition>'
    elif transition_type == "wipe":
        xml = f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="{speed}"><p:wipe dir="{direction}"/></p:transition>'
    else: # fade
        xml = f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="{speed}"><p:fade/></p:transition>'
        
    elem = parse_xml(xml)
    slide._element.append(elem)

def add_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, pt_to_in(960), pt_to_in(540))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.color.rgb = COLOR_BG
    bg.line.width = Pt(0)
    return bg

def add_header(slide, section_tag, title_text, subtitle_text):
    shapes = slide.shapes
    
    # 1. Top Logo Block: Terracotta Red Square with White 't'
    sq = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt_to_in(42), pt_to_in(25), pt_to_in(22), pt_to_in(22))
    sq.fill.solid()
    sq.fill.fore_color.rgb = COLOR_RED
    sq.line.color.rgb = COLOR_RED
    sq.line.width = Pt(0)
    p_sq = sq.text_frame.paragraphs[0]
    p_sq.text = "t"
    p_sq.font.name = FONT_FAMILY
    p_sq.font.size = Pt(13)
    p_sq.font.bold = True
    p_sq.font.color.rgb = COLOR_WHITE
    p_sq.alignment = PP_ALIGN.CENTER
    sq.text_frame.margin_top = Inches(0)
    sq.text_frame.margin_bottom = Inches(0)
    
    # Logo Text: 'timora.'
    tx_brand = shapes.add_textbox(pt_to_in(72), pt_to_in(24), pt_to_in(100), pt_to_in(25))
    tf_b = tx_brand.text_frame
    tf_b.margin_top = Inches(0)
    tf_b.margin_left = Inches(0)
    p_b = tf_b.paragraphs[0]
    p_b.text = "timora."
    p_b.font.name = FONT_FAMILY
    p_b.font.size = Pt(14)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_DARK
    
    # Top Right: 'PROJECT PRESENTATION'
    tx_proj = shapes.add_textbox(pt_to_in(750), pt_to_in(26), pt_to_in(168), pt_to_in(25))
    tf_p = tx_proj.text_frame
    tf_p.margin_top = Inches(0)
    tf_p.margin_right = Inches(0)
    p_p = tf_p.paragraphs[0]
    p_p.text = "PROJECT PRESENTATION"
    p_p.font.name = FONT_FAMILY
    p_p.font.size = Pt(8.5)
    p_p.font.bold = True
    p_p.font.color.rgb = COLOR_MUTED
    p_p.alignment = PP_ALIGN.RIGHT
    
    # Section Tag (Terracotta Red)
    tx_sec = shapes.add_textbox(pt_to_in(42), pt_to_in(64), pt_to_in(500), pt_to_in(20))
    tf_sec = tx_sec.text_frame
    tf_sec.margin_top = Inches(0)
    tf_sec.margin_left = Inches(0)
    p_sec = tf_sec.paragraphs[0]
    p_sec.text = section_tag
    p_sec.font.name = FONT_FAMILY
    p_sec.font.size = Pt(10)
    p_sec.font.bold = True
    p_sec.font.color.rgb = COLOR_RED
    
    # Main Slide Title
    tx_t = shapes.add_textbox(pt_to_in(42), pt_to_in(82), pt_to_in(700), pt_to_in(42))
    tf_t = tx_t.text_frame
    tf_t.margin_top = Inches(0)
    tf_t.margin_left = Inches(0)
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_FAMILY
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_DARK
    
    # Subtitle
    if subtitle_text:
        tx_sub = shapes.add_textbox(pt_to_in(42), pt_to_in(124), pt_to_in(750), pt_to_in(24))
        tf_s = tx_sub.text_frame
        tf_s.margin_top = Inches(0)
        tf_s.margin_left = Inches(0)
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(12)
        p_s.font.color.rgb = COLOR_MUTED

def add_footer(slide, current_idx):
    shapes = slide.shapes
    
    # Thin divider line at y=505 pt
    line = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(42), pt_to_in(505), pt_to_in(876), pt_to_in(1))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.color.rgb = COLOR_BORDER
    line.line.width = Pt(0)
    
    # Left text: 'TIMORA  /  LOCAL EVALUATION BUILD'
    tx_l = shapes.add_textbox(pt_to_in(42), pt_to_in(512), pt_to_in(300), pt_to_in(20))
    tf_l = tx_l.text_frame
    tf_l.margin_top = Inches(0)
    tf_l.margin_left = Inches(0)
    p_l = tf_l.paragraphs[0]
    p_l.text = "TIMORA  /  LOCAL EVALUATION BUILD"
    p_l.font.name = FONT_FAMILY
    p_l.font.size = Pt(8)
    p_l.font.bold = True
    p_l.font.color.rgb = COLOR_MUTED
    
    # Right text: '0X / 06'
    tx_r = shapes.add_textbox(pt_to_in(850), pt_to_in(512), pt_to_in(68), pt_to_in(20))
    tf_r = tx_r.text_frame
    tf_r.margin_top = Inches(0)
    tf_r.margin_right = Inches(0)
    p_r = tf_r.paragraphs[0]
    p_r.text = f"{current_idx:02d} / 06"
    p_r.font.name = FONT_FAMILY
    p_r.font.size = Pt(8)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_MUTED
    p_r.alignment = PP_ALIGN.RIGHT

def add_pill(slide, x, y, w, h, text, text_color=COLOR_PILL_TEXT, bg_color=COLOR_PILL_BG, font_size=9):
    shapes = slide.shapes
    pill = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt_to_in(x), pt_to_in(y), pt_to_in(w), pt_to_in(h))
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg_color
    pill.line.color.rgb = bg_color
    pill.line.width = Pt(0)
    p = pill.text_frame.paragraphs[0]
    p.text = text
    p.font.name = FONT_FAMILY
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    pill.text_frame.margin_top = Inches(0.02)
    pill.text_frame.margin_bottom = Inches(0.02)
    return pill

# -------------------------------------------------------------------
# BUILD EACH SLIDE TO MATCH THE PDF EXACTLY
# -------------------------------------------------------------------

def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    set_slide_transition(slide, "fade", speed="med")
    
    add_header(
        slide,
        "01  THE IDEA",
        "A clearer academic day",
        "One workspace for the people, places, and decisions behind a timetable."
    )
    
    shapes = slide.shapes
    
    # Sage pill
    add_pill(slide, 42, 171, 132.5, 22, "ACADEMIC OPERATIONS", COLOR_PILL_TEXT, COLOR_PILL_BG, 8.5)
    
    # Large Headline
    tx_hl = shapes.add_textbox(pt_to_in(42), pt_to_in(204), pt_to_in(380), pt_to_in(65))
    tf_hl = tx_hl.text_frame
    tf_hl.word_wrap = True
    tf_hl.margin_left = Inches(0)
    p_hl = tf_hl.paragraphs[0]
    p_hl.text = "Every timetable change has a\nripple effect."
    p_hl.font.name = FONT_FAMILY
    p_hl.font.size = Pt(22)
    p_hl.font.bold = True
    p_hl.font.color.rgb = COLOR_DARK
    
    # Body Paragraph
    tx_body = shapes.add_textbox(pt_to_in(42), pt_to_in(328), pt_to_in(390), pt_to_in(80))
    tf_b = tx_body.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0)
    p_b = tf_b.paragraphs[0]
    p_b.text = "Timora connects scheduling, resource allocation, and\nconflict review so teams can make changes with\ncontext and confidence."
    p_b.font.name = FONT_FAMILY
    p_b.font.size = Pt(13.5)
    p_b.font.color.rgb = COLOR_MUTED
    
    # Bottom Label Tag (Terracotta Red)
    tx_lbl = shapes.add_textbox(pt_to_in(42), pt_to_in(454), pt_to_in(250), pt_to_in(20))
    tf_lbl = tx_lbl.text_frame
    tf_lbl.margin_left = Inches(0)
    p_lbl = tf_lbl.paragraphs[0]
    p_lbl.text = "REAL PRODUCT SCREEN  /  LOCAL BUILD"
    p_lbl.font.name = FONT_FAMILY
    p_lbl.font.size = Pt(9)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = COLOR_RED
    
    # Right Image Frame
    # Offset shadow card: [485, 185, 435, 277]
    card_sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(485), pt_to_in(185), pt_to_in(435), pt_to_in(277))
    card_sh.fill.solid()
    card_sh.fill.fore_color.rgb = COLOR_SHADOW
    card_sh.line.color.rgb = COLOR_SHADOW
    card_sh.line.width = Pt(0)
    
    # Image: page_1_img_1.png at [480, 180, 435, 277]
    img_path = "presentation/extracted_images/page_1_img_1.png"
    if os.path.exists(img_path):
        shapes.add_picture(img_path, pt_to_in(480), pt_to_in(180), pt_to_in(435), pt_to_in(277))
        
    card_bd = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(480), pt_to_in(180), pt_to_in(435), pt_to_in(277))
    card_bd.fill.background()
    card_bd.line.color.rgb = COLOR_BORDER
    card_bd.line.width = Pt(1)
    
    add_footer(slide, 1)

def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    set_slide_transition(slide, "push", speed="med", direction="l")
    
    add_header(
        slide,
        "02  PRODUCT OVERVIEW",
        "The week, in one view",
        "Operational signals and the workflows behind them share the same scheduling records."
    )
    
    shapes = slide.shapes
    
    # Left Image Frame at [42, 156, 625, 303]
    card_sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(47), pt_to_in(161), pt_to_in(625), pt_to_in(303))
    card_sh.fill.solid()
    card_sh.fill.fore_color.rgb = COLOR_SHADOW
    card_sh.line.color.rgb = COLOR_SHADOW
    card_sh.line.width = Pt(0)
    
    img_path = "presentation/extracted_images/page_2_img_1.png"
    if os.path.exists(img_path):
        shapes.add_picture(img_path, pt_to_in(42), pt_to_in(156), pt_to_in(625), pt_to_in(303))
        
    card_bd = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(42), pt_to_in(156), pt_to_in(625), pt_to_in(303))
    card_bd.fill.background()
    card_bd.line.color.rgb = COLOR_BORDER
    card_bd.line.width = Pt(1)
    
    # Right 3 Cards at x=692, w=226, h=75
    right_cards = [
        (168, "See the state", "Sessions, conflicts, rooms, approvals, and\nutilisation together."),
        (264, "Follow the issue", "Move from an overview signal to a\nfocused workflow."),
        (360, "Work across teams", "Faculty, rooms, requests, analytics, and\nhistory stay connected."),
    ]
    
    for y_pos, title, desc in right_cards:
        # Card background
        c_bg = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt_to_in(692), pt_to_in(y_pos), pt_to_in(226), pt_to_in(75))
        c_bg.fill.solid()
        c_bg.fill.fore_color.rgb = COLOR_CARD_BG
        c_bg.line.color.rgb = COLOR_CARD_BG
        c_bg.line.width = Pt(0)
        
        # Text
        tx = shapes.add_textbox(pt_to_in(706), pt_to_in(y_pos + 12), pt_to_in(204), pt_to_in(55))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_left = Inches(0)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_DARK
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_MUTED
        
    add_footer(slide, 2)

def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    set_slide_transition(slide, "push", speed="med", direction="l")
    
    add_header(
        slide,
        "03  SMART TIMETABLE MANAGEMENT",
        "Plan. Check. Resolve.",
        "A coordinator can inspect the week, stage a change, and validate before saving."
    )
    
    shapes = slide.shapes
    
    # Left Image: Timetable Studio at [42, 156, 430, 239]
    card_sh1 = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(47), pt_to_in(161), pt_to_in(430), pt_to_in(239))
    card_sh1.fill.solid()
    card_sh1.fill.fore_color.rgb = COLOR_SHADOW
    card_sh1.line.color.rgb = COLOR_SHADOW
    card_sh1.line.width = Pt(0)
    
    img1_path = "presentation/extracted_images/page_3_img_1.png"
    if os.path.exists(img1_path):
        shapes.add_picture(img1_path, pt_to_in(42), pt_to_in(156), pt_to_in(430), pt_to_in(239))
        
    card_bd1 = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(42), pt_to_in(156), pt_to_in(430), pt_to_in(239))
    card_bd1.fill.background()
    card_bd1.line.color.rgb = COLOR_BORDER
    card_bd1.line.width = Pt(1)
    
    # Left Labels
    tx_l1 = shapes.add_textbox(pt_to_in(42), pt_to_in(405), pt_to_in(300), pt_to_in(18))
    tf_l1 = tx_l1.text_frame
    tf_l1.margin_left = Inches(0)
    p_l1 = tf_l1.paragraphs[0]
    p_l1.text = "TIMETABLE STUDIO"
    p_l1.font.name = FONT_FAMILY
    p_l1.font.size = Pt(9)
    p_l1.font.bold = True
    p_l1.font.color.rgb = COLOR_RED
    
    tx_l2 = shapes.add_textbox(pt_to_in(42), pt_to_in(423), pt_to_in(380), pt_to_in(20))
    tf_l2 = tx_l2.text_frame
    tf_l2.margin_left = Inches(0)
    p_l2 = tf_l2.paragraphs[0]
    p_l2.text = "Day-by-day sessions, course, section, room, and teacher."
    p_l2.font.name = FONT_FAMILY
    p_l2.font.size = Pt(10.5)
    p_l2.font.color.rgb = COLOR_MUTED
    
    # Right Image: Conflict Center at [488, 156, 430, 239]
    card_sh2 = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(493), pt_to_in(161), pt_to_in(430), pt_to_in(239))
    card_sh2.fill.solid()
    card_sh2.fill.fore_color.rgb = COLOR_SHADOW
    card_sh2.line.color.rgb = COLOR_SHADOW
    card_sh2.line.width = Pt(0)
    
    img2_path = "presentation/extracted_images/page_3_img_2.png"
    if os.path.exists(img2_path):
        shapes.add_picture(img2_path, pt_to_in(488), pt_to_in(156), pt_to_in(430), pt_to_in(239))
        
    card_bd2 = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(488), pt_to_in(156), pt_to_in(430), pt_to_in(239))
    card_bd2.fill.background()
    card_bd2.line.color.rgb = COLOR_BORDER
    card_bd2.line.width = Pt(1)
    
    # Right Labels
    tx_r1 = shapes.add_textbox(pt_to_in(488), pt_to_in(405), pt_to_in(300), pt_to_in(18))
    tf_r1 = tx_r1.text_frame
    tf_r1.margin_left = Inches(0)
    p_r1 = tf_r1.paragraphs[0]
    p_r1.text = "CONFLICT CENTER"
    p_r1.font.name = FONT_FAMILY
    p_r1.font.size = Pt(9)
    p_r1.font.bold = True
    p_r1.font.color.rgb = COLOR_RED
    
    tx_r2 = shapes.add_textbox(pt_to_in(488), pt_to_in(423), pt_to_in(380), pt_to_in(20))
    tf_r2 = tx_r2.text_frame
    tf_r2.margin_left = Inches(0)
    p_r2 = tf_r2.paragraphs[0]
    p_r2.text = "Room and faculty clashes with reviewable alternatives."
    p_r2.font.name = FONT_FAMILY
    p_r2.font.size = Pt(10.5)
    p_r2.font.color.rgb = COLOR_MUTED
    
    # Bottom Sage Pill
    add_pill(slide, 42, 467, 309.8, 22, "CSV + XLSX IMPORT: MATCH  /  REVIEW  /  VALIDATE  /  CONFIRM", COLOR_PILL_TEXT, COLOR_PILL_BG, 8.5)
    
    add_footer(slide, 3)

def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    set_slide_transition(slide, "push", speed="med", direction="l")
    
    add_header(
        slide,
        "04  TIMO AI COPILOT",
        "Answers tied to records",
        "A real CS-401 lookup from the local Timora workspace."
    )
    
    shapes = slide.shapes
    
    # Left Content
    tx_lbl = shapes.add_textbox(pt_to_in(42), pt_to_in(170), pt_to_in(300), pt_to_in(18))
    tf_lbl = tx_lbl.text_frame
    tf_lbl.margin_left = Inches(0)
    p_lbl = tf_lbl.paragraphs[0]
    p_lbl.text = "WHAT IT DOES TODAY"
    p_lbl.font.name = FONT_FAMILY
    p_lbl.font.size = Pt(9)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = COLOR_RED
    
    tx_sub = shapes.add_textbox(pt_to_in(42), pt_to_in(188), pt_to_in(320), pt_to_in(30))
    tf_sub = tx_sub.text_frame
    tf_sub.margin_left = Inches(0)
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Ask in natural language."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(17)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_DARK
    
    tx_p1 = shapes.add_textbox(pt_to_in(42), pt_to_in(256), pt_to_in(340), pt_to_in(60))
    tf_p1 = tx_p1.text_frame
    tf_p1.word_wrap = True
    tf_p1.margin_left = Inches(0)
    p_p1 = tf_p1.paragraphs[0]
    p_p1.text = "Timo retrieves timetable, availability, conflict, workload,\nsyllabus, and version data through bounded tools."
    p_p1.font.name = FONT_FAMILY
    p_p1.font.size = Pt(12)
    p_p1.font.color.rgb = COLOR_MUTED
    
    tx_p2 = shapes.add_textbox(pt_to_in(42), pt_to_in(338), pt_to_in(340), pt_to_in(75))
    tf_p2 = tx_p2.text_frame
    tf_p2.word_wrap = True
    tf_p2.margin_left = Inches(0)
    p_p2 = tf_p2.paragraphs[0]
    p_p2.text = "Source records and follow-up suggestions make answers\ntraceable. Supported changes are previewed first and\nrequire user confirmation."
    p_p2.font.name = FONT_FAMILY
    p_p2.font.size = Pt(12)
    p_p2.font.color.rgb = COLOR_MUTED
    
    # Pill
    add_pill(slide, 42, 451, 164.5, 22, "READ  >  PROPOSE  >  CONFIRM", COLOR_PILL_TEXT, COLOR_PILL_BG, 8.5)
    
    # Right Image Frame: [399, 153, 519, 312]
    card_sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(404), pt_to_in(158), pt_to_in(519), pt_to_in(312))
    card_sh.fill.solid()
    card_sh.fill.fore_color.rgb = COLOR_SHADOW
    card_sh.line.color.rgb = COLOR_SHADOW
    card_sh.line.width = Pt(0)
    
    img_path = "presentation/extracted_images/page_4_img_1.png"
    if os.path.exists(img_path):
        shapes.add_picture(img_path, pt_to_in(399), pt_to_in(153), pt_to_in(519), pt_to_in(312))
        
    card_bd = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(399), pt_to_in(153), pt_to_in(519), pt_to_in(312))
    card_bd.fill.background()
    card_bd.line.color.rgb = COLOR_BORDER
    card_bd.line.width = Pt(1)
    
    add_footer(slide, 4)

def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    set_slide_transition(slide, "push", speed="med", direction="l")
    
    add_header(
        slide,
        "05  TECHNOLOGY & ARCHITECTURE",
        "A shared source of truth",
        "The current code separates the interface, scheduling logic, data access, and AI service."
    )
    
    shapes = slide.shapes
    
    # 4 Architecture Flow Cards at y=184, h=122
    cards = [
        (42, 184, "INTERFACE", "Next.js 16 + React 19", "TypeScript + Tailwind"),
        (269, 184, "SCHEDULING LAYER", "Repository + constraints", "Plans, imports, requests"),
        (496, 184, "DATA PATH", "Local evaluation store", "Supabase live adapter"),
        (723, 195, "TIMO SERVICE", "Server-side OpenRouter", "Validated, scoped tools"),
    ]
    
    for x_pos, w_pos, tag, line1, line2 in cards:
        card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt_to_in(x_pos), pt_to_in(184), pt_to_in(w_pos), pt_to_in(122))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)
        
        tx = shapes.add_textbox(pt_to_in(x_pos + 16), pt_to_in(200), pt_to_in(w_pos - 32), pt_to_in(90))
        tf = tx.text_frame
        tf.margin_top = Inches(0)
        tf.margin_left = Inches(0)
        
        p1 = tf.paragraphs[0]
        p1.text = tag
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_RED
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{line1}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_DARK
        
        p3 = tf.add_paragraph()
        p3.text = line2
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_DARK
        
    # Flow Connecting Arrows (Terracotta Red)
    arrows_x = [235, 462, 689]
    for ax in arrows_x:
        # horizontal line
        arr = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(ax), pt_to_in(244), pt_to_in(25), pt_to_in(2))
        arr.fill.solid()
        arr.fill.fore_color.rgb = COLOR_RED
        arr.line.color.rgb = COLOR_RED
        arr.line.width = Pt(0)
        
    # Two Text Columns below
    # Left column: VALIDATION BEFORE WRITE
    tx_v1 = shapes.add_textbox(pt_to_in(42), pt_to_in(340), pt_to_in(300), pt_to_in(18))
    tf_v1 = tx_v1.text_frame
    tf_v1.margin_left = Inches(0)
    p_v1 = tf_v1.paragraphs[0]
    p_v1.text = "VALIDATION BEFORE WRITE"
    p_v1.font.name = FONT_FAMILY
    p_v1.font.size = Pt(9)
    p_v1.font.bold = True
    p_v1.font.color.rgb = COLOR_RED
    
    tx_v2 = shapes.add_textbox(pt_to_in(42), pt_to_in(362), pt_to_in(440), pt_to_in(45))
    tf_v2 = tx_v2.text_frame
    tf_v2.word_wrap = True
    tf_v2.margin_left = Inches(0)
    p_v2 = tf_v2.paragraphs[0]
    p_v2.text = "Constraint checks, role checks, confirmation, audit, and read-after-write verification protect local workflows."
    p_v2.font.name = FONT_FAMILY
    p_v2.font.size = Pt(11.5)
    p_v2.font.color.rgb = COLOR_MUTED
    
    # Right column: LIVE DEPLOYMENT GATE
    tx_g1 = shapes.add_textbox(pt_to_in(506), pt_to_in(340), pt_to_in(300), pt_to_in(18))
    tf_g1 = tx_g1.text_frame
    tf_g1.margin_left = Inches(0)
    p_g1 = tf_g1.paragraphs[0]
    p_g1.text = "LIVE DEPLOYMENT GATE"
    p_g1.font.name = FONT_FAMILY
    p_g1.font.size = Pt(9)
    p_g1.font.bold = True
    p_g1.font.color.rgb = COLOR_RED
    
    tx_g2 = shapes.add_textbox(pt_to_in(506), pt_to_in(362), pt_to_in(412), pt_to_in(45))
    tf_g2 = tx_g2.text_frame
    tf_g2.word_wrap = True
    tf_g2.margin_left = Inches(0)
    p_g2 = tf_g2.paragraphs[0]
    p_g2.text = "Supabase row policies, transactional writes, and live role tests still require completion."
    p_g2.font.name = FONT_FAMILY
    p_g2.font.size = Pt(11.5)
    p_g2.font.color.rgb = COLOR_MUTED
    
    # Bottom Pill
    add_pill(slide, 42, 456, 198.5, 22, "LOCAL SCREENS SHOWN IN THIS DECK", COLOR_PILL_TEXT, COLOR_PILL_BG, 8.5)
    
    add_footer(slide, 5)

def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    set_slide_transition(slide, "fade", speed="med")
    
    add_header(
        slide,
        "06  IMPACT & NEXT STEPS",
        "Make every change trustworthy",
        "Timora connects insight, planning, and review in a single academic workflow."
    )
    
    shapes = slide.shapes
    
    # Left Image Frame at [42, 156, 446, 271]
    card_sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(47), pt_to_in(161), pt_to_in(446), pt_to_in(271))
    card_sh.fill.solid()
    card_sh.fill.fore_color.rgb = COLOR_SHADOW
    card_sh.line.color.rgb = COLOR_SHADOW
    card_sh.line.width = Pt(0)
    
    img_path = "presentation/extracted_images/page_6_img_1.png"
    if os.path.exists(img_path):
        shapes.add_picture(img_path, pt_to_in(42), pt_to_in(156), pt_to_in(446), pt_to_in(271))
        
    card_bd = shapes.add_shape(MSO_SHAPE.RECTANGLE, pt_to_in(42), pt_to_in(156), pt_to_in(446), pt_to_in(271))
    card_bd.fill.background()
    card_bd.line.color.rgb = COLOR_BORDER
    card_bd.line.width = Pt(1)
    
    # Right Side Content
    # PRACTICAL VALUE
    tx_pv = shapes.add_textbox(pt_to_in(516), pt_to_in(164), pt_to_in(300), pt_to_in(18))
    tf_pv = tx_pv.text_frame
    tf_pv.margin_left = Inches(0)
    p_pv = tf_pv.paragraphs[0]
    p_pv.text = "PRACTICAL VALUE"
    p_pv.font.name = FONT_FAMILY
    p_pv.font.size = Pt(9)
    p_pv.font.bold = True
    p_pv.font.color.rgb = COLOR_RED
    
    tx_pvt = shapes.add_textbox(pt_to_in(516), pt_to_in(184), pt_to_in(400), pt_to_in(50))
    tf_pvt = tx_pvt.text_frame
    tf_pvt.word_wrap = True
    tf_pvt.margin_left = Inches(0)
    p_pvt = tf_pvt.paragraphs[0]
    p_pvt.text = "See conflicts earlier. Compare options. Keep\na trace of decisions."
    p_pvt.font.name = FONT_FAMILY
    p_pvt.font.size = Pt(17)
    p_pvt.font.bold = True
    p_pvt.font.color.rgb = COLOR_DARK
    
    # ROADMAP TO INSTITUTIONAL LAUNCH
    tx_rm = shapes.add_textbox(pt_to_in(516), pt_to_in(276), pt_to_in(300), pt_to_in(18))
    tf_rm = tx_rm.text_frame
    tf_rm.margin_left = Inches(0)
    p_rm = tf_rm.paragraphs[0]
    p_rm.text = "ROADMAP TO INSTITUTIONAL LAUNCH"
    p_rm.font.name = FONT_FAMILY
    p_rm.font.size = Pt(9)
    p_rm.font.bold = True
    p_rm.font.color.rgb = COLOR_RED
    
    # 4 Numbered Steps
    steps = [
        (299, "1", "Scoped Supabase row policies"),
        (331, "2", "Transactional, audited live writes"),
        (363, "3", "Complete exchange consent and application"),
        (395, "4", "Live permission, recovery, and accessibility tests"),
    ]
    
    for y_pos, num, text in steps:
        # Sage Number Square [514, y, 24, 24]
        sq_n = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt_to_in(514), pt_to_in(y_pos), pt_to_in(24), pt_to_in(24))
        sq_n.fill.solid()
        sq_n.fill.fore_color.rgb = COLOR_PILL_BG
        sq_n.line.color.rgb = COLOR_PILL_BG
        sq_n.line.width = Pt(0)
        p_n = sq_n.text_frame.paragraphs[0]
        p_n.text = num
        p_n.font.name = FONT_FAMILY
        p_n.font.size = Pt(10)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_PILL_TEXT
        p_n.alignment = PP_ALIGN.CENTER
        sq_n.text_frame.margin_top = Inches(0)
        sq_n.text_frame.margin_bottom = Inches(0)
        
        # Step Label
        tx_s = shapes.add_textbox(pt_to_in(547), pt_to_in(y_pos + 2), pt_to_in(370), pt_to_in(22))
        tf_s = tx_s.text_frame
        tf_s.margin_left = Inches(0)
        tf_s.margin_top = Inches(0)
        p_s = tf_s.paragraphs[0]
        p_s.text = text
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = COLOR_DARK
        
    # Bottom Forest Charcoal Banner: [42, 444, 876, 37]
    banner = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt_to_in(42), pt_to_in(444), pt_to_in(876), pt_to_in(37))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_BANNER_BG
    banner.line.color.rgb = COLOR_BANNER_BG
    banner.line.width = Pt(0)
    p_ban = banner.text_frame.paragraphs[0]
    p_ban.text = "The right people. The right place. The right time. A change you can explain."
    p_ban.font.name = FONT_FAMILY
    p_ban.font.size = Pt(14)
    p_ban.font.bold = True
    p_ban.font.color.rgb = COLOR_WHITE
    p_ban.alignment = PP_ALIGN.CENTER
    banner.text_frame.margin_top = Inches(0.04)
    
    add_footer(slide, 6)

def main():
    print("Building Timora Project Presentation PPTX from original PDF structure...")
    prs = Presentation()
    prs.slide_width = pt_to_in(960)
    prs.slide_height = pt_to_in(540)
    
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    
    out_path = "presentation/Timora_Project_Presentation.pptx"
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    main()

